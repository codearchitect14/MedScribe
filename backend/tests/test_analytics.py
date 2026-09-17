"""Tests for the Phase 7 analytics rollups, reimbursement rate table, and
CSV/Excel export.

Drives an encounter through the full lifecycle (transcribed -> ... ->
billed) using the same mocked-adapter pattern as tests/test_billing.py,
sets a reimbursement rate, triggers an on-demand rollup, and checks the
resulting numbers against what was actually created - not just that the
endpoints return 200.
"""

import csv
import io
import json
import re
import uuid
from datetime import date

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.security import hash_password
from app.llm import gateway
from app.llm.adapters.base import ProviderResponse
from app.main import app
from app.models.organization import Organization
from app.models.user import User, UserRole

settings = get_settings()

def _make_soap_json() -> str:
    # A unique marker per instantiation, so this test's LLM prompts never
    # collide with the same-shaped calls other test files make: the
    # request-level cache (app/llm/cache.py) is keyed on prompt content in
    # the same long-lived dev Redis instance, and a cache hit skips writing
    # an llm_usage_logs row entirely, which would starve the llm_usage
    # rollup this test checks.
    return json.dumps(
        {
            "subjective": "s",
            "objective": "o",
            "assessment": f"Type 2 diabetes mellitus with hyperglycemia. [{uuid.uuid4()}]",
            "plan": "p",
        }
    )


class ScriptedAdapter:
    name = "groq"

    def __init__(self):
        self.soap_json = _make_soap_json()

    async def complete(self, *, system_prompt, user_prompt, max_tokens, json_mode=True):
        if max_tokens == 500:
            content = self.soap_json
        elif max_tokens == 200:
            match = re.search(r"-\s*(\S+):", user_prompt)
            code = match.group(1) if match else "E11.9"
            content = json.dumps(
                {"selected_codes": [{"code": code, "justification": "matches", "confidence": 0.9}]}
            )
        else:
            content = json.dumps({"diagnosis_summary": "n/a", "next_appointment_guidance": "n/a"})
        return ProviderResponse(content=content, tokens_input=120, tokens_output=40, model="mock-model")


def _require_db():
    engine = create_engine(settings.sync_database_url)
    try:
        conn = engine.connect()
        conn.close()
    except OperationalError:
        pytest.skip("database not reachable; start docker compose postgres service first")


@pytest.fixture
def scripted_adapter(monkeypatch):
    adapter = ScriptedAdapter()
    monkeypatch.setitem(gateway.ADAPTERS, "groq", adapter)
    monkeypatch.setitem(gateway.ADAPTERS, "gemini", adapter)
    return adapter


@pytest.fixture
def clinician_admin_coder():
    _require_db()
    suffix = uuid.uuid4().hex[:8]
    with Session(create_engine(settings.sync_database_url)) as session:
        org = Organization(name=f"Analytics Test Org {suffix}")
        session.add(org)
        session.flush()

        users = {}
        for role in (UserRole.clinician, UserRole.admin, UserRole.coder):
            email = f"{role.value}.{suffix}@example-medscribe.com"
            user = User(
                email=email,
                hashed_password=hash_password("TestPass123!"),
                full_name=f"Test {role.value}",
                role=role,
                organization_id=org.id,
                is_active=True,
            )
            session.add(user)
            users[role.value] = {"email": email, "password": "TestPass123!", "id": None}

        session.commit()
        org_id = org.id

    return {"org_id": org_id, "users": users}


async def _login(client: AsyncClient, email: str, password: str) -> str:
    response = await client.post("/auth/login", data={"username": email, "password": password})
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


async def _drive_encounter_to_billed(client: AsyncClient, headers: dict) -> str:
    patient = await client.post(
        "/patients", json={"first_name": "Analytics", "last_name": "Test"}, headers=headers
    )
    patient_id = patient.json()["id"]

    # The transcript must be unique per call: the LLM gateway's request-level
    # cache (app/llm/cache.py) is keyed on the input prompt in the same
    # long-lived dev Redis instance, and a cache hit skips writing an
    # llm_usage_logs row entirely, which would starve the llm_usage rollup.
    encounter = await client.post(
        "/encounters",
        json={
            "patient_id": patient_id,
            "raw_transcript": f"Doctor: hi. Patient: sugar is high. [{uuid.uuid4()}]",
        },
        headers=headers,
    )
    encounter_id = encounter.json()["id"]

    await client.post(f"/encounters/{encounter_id}/soap-note", headers=headers)
    await client.post(f"/encounters/{encounter_id}/soap-note/review", headers=headers)
    await client.post(f"/encounters/{encounter_id}/soap-note/finalize", headers=headers)
    await client.post(f"/encounters/{encounter_id}/codes", headers=headers)
    return encounter_id


@pytest.mark.asyncio(loop_scope="session")
async def test_reimbursement_rate_crud_and_rbac(clinician_admin_coder):
    users = clinician_admin_coder["users"]
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        admin_token = await _login(client, users["admin"]["email"], users["admin"]["password"])
        coder_token = await _login(client, users["coder"]["email"], users["coder"]["password"])
        clinician_token = await _login(client, users["clinician"]["email"], users["clinician"]["password"])
        admin_headers = {"Authorization": f"Bearer {admin_token}"}
        coder_headers = {"Authorization": f"Bearer {coder_token}"}
        clinician_headers = {"Authorization": f"Bearer {clinician_token}"}

        # Clinicians cannot manage rates at all.
        forbidden_read = await client.get("/reimbursement-rates", headers=clinician_headers)
        assert forbidden_read.status_code == 403

        # Coders can read but not write.
        forbidden_write = await client.put(
            "/reimbursement-rates",
            json={"code_type": "icd10", "code": "E11.9", "rate": 150.0},
            headers=coder_headers,
        )
        assert forbidden_write.status_code == 403

        created = await client.put(
            "/reimbursement-rates",
            json={"code_type": "icd10", "code": "E11.9", "rate": 150.0},
            headers=admin_headers,
        )
        assert created.status_code == 200, created.text
        assert created.json()["rate"] == 150.0
        rate_id = created.json()["id"]

        # Re-upserting the same code replaces the rate rather than duplicating.
        updated = await client.put(
            "/reimbursement-rates",
            json={"code_type": "icd10", "code": "E11.9", "rate": 175.5},
            headers=admin_headers,
        )
        assert updated.status_code == 200
        assert updated.json()["id"] == rate_id
        assert updated.json()["rate"] == 175.5

        listing = await client.get("/reimbursement-rates", headers=coder_headers)
        assert listing.status_code == 200
        assert listing.json()["total"] == 1

        deleted = await client.delete(f"/reimbursement-rates/{rate_id}", headers=admin_headers)
        assert deleted.status_code == 204

        empty_listing = await client.get("/reimbursement-rates", headers=admin_headers)
        assert empty_listing.json()["total"] == 0


@pytest.mark.asyncio(loop_scope="session")
async def test_analytics_rollup_end_to_end(clinician_admin_coder, scripted_adapter):
    users = clinician_admin_coder["users"]
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        clinician_token = await _login(client, users["clinician"]["email"], users["clinician"]["password"])
        admin_token = await _login(client, users["admin"]["email"], users["admin"]["password"])
        coder_token = await _login(client, users["coder"]["email"], users["coder"]["password"])
        clinician_headers = {"Authorization": f"Bearer {clinician_token}"}
        admin_headers = {"Authorization": f"Bearer {admin_token}"}
        coder_headers = {"Authorization": f"Bearer {coder_token}"}

        me = await client.get("/auth/me", headers=clinician_headers)
        clinician_id = me.json()["id"]

        encounter_id = await _drive_encounter_to_billed(client, clinician_headers)

        # The ScriptedAdapter picks the first real pgvector-retrieved
        # candidate; which exact ICD-10 code that is depends on the real
        # embedding similarity ranking, not something to hardcode here.
        detail = await client.get(f"/encounters/{encounter_id}", headers=coder_headers)
        suggestion = detail.json()["code_suggestions"][0]
        suggestion_id = suggestion["id"]
        accepted_code = suggestion["code"]
        assert suggestion["code_type"] == "icd10"

        rate = await client.put(
            "/reimbursement-rates",
            json={"code_type": "icd10", "code": accepted_code, "rate": 150.0},
            headers=admin_headers,
        )
        assert rate.status_code == 200

        # Accept the code and create a billing record so estimated
        # reimbursement and the finalize->billed turnaround have data.
        await client.patch(
            f"/encounters/{encounter_id}/codes/{suggestion_id}",
            json={"accepted": True},
            headers=coder_headers,
        )
        billing = await client.post(
            f"/encounters/{encounter_id}/billing-record",
            json={"codes_applied": [accepted_code], "payer": "Acme"},
            headers=coder_headers,
        )
        assert billing.status_code == 201, billing.text

        # Clinicians cannot trigger or view analytics.
        assert (await client.post("/analytics/rollup", json={}, headers=clinician_headers)).status_code == 403

        today = date.today().isoformat()
        triggered = await client.post("/analytics/rollup", json={"period": today}, headers=admin_headers)
        assert triggered.status_code == 200, triggered.text
        assert set(triggered.json()["rollup_types"]) == {
            "encounter_volume",
            "coding_mix",
            "estimated_reimbursement",
            "turnaround_time",
            "llm_usage",
            "clinician_productivity",
        }

        volume = await client.get(
            "/analytics/encounter_volume",
            params={"start_date": today, "end_date": today},
            headers=admin_headers,
        )
        assert volume.status_code == 200
        volume_rows = volume.json()["rollups"][0]["dimensions"]
        clinician_row = next(r for r in volume_rows if r["clinician_id"] == clinician_id)
        assert clinician_row["count"] == 1
        total_row = next(r for r in volume_rows if r["clinician_id"] is None)
        assert total_row["count"] == 1

        coding_mix = await client.get(
            "/analytics/coding_mix", params={"start_date": today, "end_date": today}, headers=admin_headers
        )
        coding_rows = coding_mix.json()["rollups"][0]["dimensions"]
        assert any(r["code_type"] == "icd10" and r["count"] == 1 for r in coding_rows)

        reimbursement = await client.get(
            "/analytics/estimated_reimbursement",
            params={"start_date": today, "end_date": today},
            headers=admin_headers,
        )
        reimb_row = reimbursement.json()["rollups"][0]["dimensions"][0]
        assert reimb_row["estimated_reimbursement"] == 150.0
        assert reimb_row["codes_matched"] == 1

        turnaround = await client.get(
            "/analytics/turnaround_time",
            params={"start_date": today, "end_date": today},
            headers=admin_headers,
        )
        turnaround_row = turnaround.json()["rollups"][0]["dimensions"][0]
        assert turnaround_row["finalized_count"] == 1
        assert turnaround_row["avg_hours_transcript_to_finalize"] >= 0
        assert turnaround_row["billed_count"] == 1
        assert turnaround_row["avg_hours_finalize_to_billed"] >= 0

        llm_usage = await client.get(
            "/analytics/llm_usage", params={"start_date": today, "end_date": today}, headers=admin_headers
        )
        llm_rows = llm_usage.json()["rollups"][0]["dimensions"]
        assert any(r["provider"] == "groq" and r["calls"] >= 1 for r in llm_rows)

        productivity = await client.get(
            "/analytics/clinician_productivity",
            params={"start_date": today, "end_date": today},
            headers=admin_headers,
        )
        prod_row = productivity.json()["rollups"][0]["dimensions"][0]
        assert prod_row["clinician_id"] == clinician_id
        assert prod_row["encounters_finalized"] == 1

        # Recomputing is idempotent: same day, same numbers, no duplicate rows.
        await client.post("/analytics/rollup", json={"period": today}, headers=admin_headers)
        volume_again = await client.get(
            "/analytics/encounter_volume",
            params={"start_date": today, "end_date": today},
            headers=admin_headers,
        )
        assert len(volume_again.json()["rollups"]) == 1


@pytest.mark.asyncio(loop_scope="session")
async def test_analytics_export_csv_and_xlsx(clinician_admin_coder, scripted_adapter):
    users = clinician_admin_coder["users"]
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        clinician_token = await _login(client, users["clinician"]["email"], users["clinician"]["password"])
        admin_token = await _login(client, users["admin"]["email"], users["admin"]["password"])
        clinician_headers = {"Authorization": f"Bearer {clinician_token}"}
        admin_headers = {"Authorization": f"Bearer {admin_token}"}

        await _drive_encounter_to_billed(client, clinician_headers)
        today = date.today().isoformat()
        await client.post("/analytics/rollup", json={"period": today}, headers=admin_headers)

        csv_response = await client.get(
            "/analytics/encounter_volume/export",
            params={"start_date": today, "end_date": today, "format": "csv"},
            headers=admin_headers,
        )
        assert csv_response.status_code == 200
        assert csv_response.headers["content-type"].startswith("text/csv")
        assert "attachment" in csv_response.headers["content-disposition"]
        rows = list(csv.DictReader(io.StringIO(csv_response.text)))
        assert len(rows) >= 1
        assert "period" in rows[0]
        assert "count" in rows[0]

        xlsx_response = await client.get(
            "/analytics/encounter_volume/export",
            params={"start_date": today, "end_date": today, "format": "xlsx"},
            headers=admin_headers,
        )
        assert xlsx_response.status_code == 200
        assert "spreadsheetml" in xlsx_response.headers["content-type"]
        workbook = load_workbook(io.BytesIO(xlsx_response.content))
        sheet = workbook.active
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        assert "period" in header
        assert "count" in header


@pytest.mark.asyncio(loop_scope="session")
async def test_analytics_scoped_to_own_organization(clinician_admin_coder, scripted_adapter):
    """An org's rollups must never leak into another org's analytics view."""
    users = clinician_admin_coder["users"]
    suffix = uuid.uuid4().hex[:8]

    with Session(create_engine(settings.sync_database_url)) as session:
        other_org = Organization(name=f"Other Analytics Org {suffix}")
        session.add(other_org)
        session.flush()
        other_admin_email = f"other-admin.{suffix}@example-medscribe.com"
        other_admin = User(
            email=other_admin_email,
            hashed_password=hash_password("TestPass123!"),
            full_name="Other Org Admin",
            role=UserRole.admin,
            organization_id=other_org.id,
            is_active=True,
        )
        session.add(other_admin)
        session.commit()

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        clinician_token = await _login(client, users["clinician"]["email"], users["clinician"]["password"])
        admin_token = await _login(client, users["admin"]["email"], users["admin"]["password"])
        clinician_headers = {"Authorization": f"Bearer {clinician_token}"}
        admin_headers = {"Authorization": f"Bearer {admin_token}"}

        await _drive_encounter_to_billed(client, clinician_headers)
        today = date.today().isoformat()
        await client.post("/analytics/rollup", json={"period": today}, headers=admin_headers)

        other_token = await _login(client, other_admin_email, "TestPass123!")
        other_headers = {"Authorization": f"Bearer {other_token}"}
        other_view = await client.get(
            "/analytics/encounter_volume",
            params={"start_date": today, "end_date": today},
            headers=other_headers,
        )
        assert other_view.status_code == 200
        assert other_view.json()["rollups"] == []
